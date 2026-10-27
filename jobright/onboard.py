"""Turn an uploaded resume into a master profile and a career plan.

import_resume()  reads a PDF, Word or text resume, asks Claude to restructure
                 it into the fact-by-fact master profile (wording kept close
                 to the original), then checks every fact back against the
                 resume text and flags any it cannot find, for you to review.
plan_career()    reads the profile and proposes which kinds of role to go
                 for: level, real strengths with their depth, transferable
                 skills, and role tracks with search keywords, from core roles
                 to realistic stretch roles.
apply_plan()     writes the plan into your settings: tracks, searches for
                 your country's job sites, title exclusions for your level.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Literal

import yaml
from pydantic import BaseModel, Field

from jobright import config
from jobright.llm import LLM
from jobright.profile import Contact, Education, Fact, Link, Profile, Project, Role


# --- reading the file ----------------------------------------------------------

def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
    if suffix in (".docx",):
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return path.read_text(errors="replace")


# --- resume → profile --------------------------------------------------------------

class XProject(BaseModel):
    name: str
    summary: str
    details: list[str] = Field(default_factory=list)


class XRole(BaseModel):
    title: str
    company: str
    engagement: str | None = Field(default=None, description="e.g. 'Client project via X', 'Internship', 'Part-time'")
    start: str = Field(description="YYYY-MM; YYYY-01 if only the year is given")
    end: str | None = Field(default=None, description="YYYY-MM, or null if current")
    bullets: list[str]
    projects: list[XProject] = Field(default_factory=list)


class XEducation(BaseModel):
    degree: str
    institution: str
    location: str | None = None
    start: str | None = None
    end: str | None = None
    grade: str | None = Field(default=None, description="e.g. 'First Class Honours', '2:1', 'CGPA 8.2/10'")
    details: list[str] = Field(default_factory=list, description="modules, dissertation, awards, as written")


class XSkillGroup(BaseModel):
    name: str
    items: list[str]


class XProfile(BaseModel):
    name: str
    location: str
    email: str
    phone: str
    links: list[Link] = Field(default_factory=list)
    titles: list[str] = Field(description="Titles the person has actually held or is qualified as, most senior first")
    years_experience: str = Field(description="As the resume implies, e.g. '8+', '1', '0 (graduate)'")
    summary: list[str]
    highlights: list[str] = Field(default_factory=list)
    skills: list[XSkillGroup]
    roles: list[XRole] = Field(default_factory=list, description="Newest first; include internships and placements")
    other_projects: list[XProject] = Field(default_factory=list, description="Personal, university or open-source projects")
    education: list[XEducation]
    certifications: list[str] = Field(default_factory=list)


def _slug(s: str, taken: set[str]) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:24] or "x"
    out, n = base, 2
    while out in taken:
        out, n = f"{base}-{n}", n + 1
    taken.add(out)
    return out


def to_profile(x: XProfile) -> Profile:
    """Give every fact a stable id, the shape tailoring and verification cite."""
    taken: set[str] = set()
    roles = []
    for r in x.roles:
        rid = _slug(r.company, taken)
        prefix = re.sub(r"[^a-z]", "", rid)[:3] or "r"
        projects = []
        for p in r.projects:
            pid = _slug(p.name, taken)
            projects.append(Project(id=pid, name=p.name, summary=p.summary,
                                    details=[Fact(id=f"{pid}-{i + 1}", text=t) for i, t in enumerate(p.details)]))
        roles.append(Role(id=rid, title=r.title, company=r.company, engagement=r.engagement, start=r.start,
                          end=r.end, bullets=[Fact(id=f"{prefix}{i + 1}-{rid}", text=t) for i, t in enumerate(r.bullets)],
                          projects=projects))
    others = []
    for p in x.other_projects:
        pid = _slug(p.name, taken)
        others.append(Project(id=pid, name=p.name, summary=p.summary,
                              details=[Fact(id=f"{pid}-{i + 1}", text=t) for i, t in enumerate(p.details)]))
    highlights = [Fact(id=f"h{i + 1}", text=t) for i, t in enumerate(x.highlights)]
    highlights += [Fact(id=f"c{i + 1}", text=f"Certification: {t}") for i, t in enumerate(x.certifications)]
    education = [Education(degree=e.degree, institution=e.institution, location=e.location, end=e.end,
                           grade=e.grade, details=[Fact(id=f"ed{j + 1}-{i + 1}", text=t) for i, t in enumerate(e.details)])
                 for j, e in enumerate(x.education)]
    return Profile(
        contact=Contact(name=x.name, location=x.location, email=x.email, phone=x.phone, links=x.links),
        titles=x.titles, years_experience=x.years_experience,
        summary=[Fact(id=f"s{i + 1}", text=t) for i, t in enumerate(x.summary)],
        highlights=highlights,
        skills={g.name: g.items for g in x.skills},
        roles=roles, other_projects=others, education=education,
    )


def _words(s: str) -> set[str]:
    return {w for w in re.findall(r"[a-z0-9+#.]{3,}", s.lower())}


def unfounded(profile: Profile, resume_text: str, cutoff: float = 0.6) -> list[tuple[str, str]]:
    """Facts whose words mostly do not appear in the resume: (id, text) to review."""
    source = _words(resume_text)
    out = []
    for fid, text in profile.facts().items():
        words = _words(text.split(": ", 1)[-1])
        if words and len(words & source) / len(words) < cutoff:
            out.append((fid, text))
    return out


def import_resume(path: Path, llm: LLM | None = None) -> tuple[Profile, list[tuple[str, str]], str]:
    text = extract_text(path)
    if len(text.strip()) < 200:
        raise ValueError("Could not read enough text from that file. Try a PDF exported from Word/Google Docs, or a .docx.")
    from jobright import engine, heuristics
    if not engine.uses_ai():
        profile = to_profile(heuristics.parse_resume(text))
        return profile, unfounded(profile, text), text
    llm = llm or LLM()
    try:
        x = _structured_profile(llm, text)
    except Exception:
        if engine.current() != "local":
            raise
        x = heuristics.parse_resume(text)     # a small local model can fail on a long CV
    profile = to_profile(x)
    return profile, unfounded(profile, text), text


def _structured_profile(llm: LLM, text: str) -> XProfile:
    return llm.structured(
        purpose="onboard_profile",
        model=config.settings()["models"]["verify"],
        instructions=config.prompt("onboard_profile"),
        context="(no prior profile)",
        prompt="<resume>\n" + text.strip() + "\n</resume>",
        schema=XProfile,
        effort="high",
        max_tokens=32000,
    )


def save_profile(profile: Profile) -> Path:
    config.DATA.mkdir(parents=True, exist_ok=True)
    data = profile.model_dump(exclude_none=True)
    header = ("# Master profile: the only source of truth for every tailored resume.\n"
              "# Created from your resume by the setup wizard. Fix anything wrong here;\n"
              "# tailoring may reorder and reword these facts but never adds new ones.\n\n")
    config.PROFILE_PATH.write_text(header + yaml.safe_dump(data, sort_keys=False, allow_unicode=True, width=120))
    return config.PROFILE_PATH


# --- profile → career plan -----------------------------------------------------------

class Strength(BaseModel):
    skill: str
    depth: Literal["expert", "solid", "working", "basic"]
    evidence: str = Field(description="Which role or project shows it, briefly")


class TrackPlan(BaseModel):
    name: str = Field(description="short snake_case id, e.g. mobile, backend_cloud, graduate_software")
    label: str
    titles: list[str] = Field(description="Phrases a matching job title contains, e.g. 'Flutter', 'Graduate Software Engineer'")
    search_keywords: list[str] = Field(description="2-4 searches to run on job sites for this track")
    weight: float = Field(description="1.0 core; up to 1.1 for high-growth tracks; 0.8-0.9 for weaker fits")
    min_fit: int = Field(description="50 for core tracks, 55-62 for stretch tracks")
    stretch: bool
    why: str


class CareerPlan(BaseModel):
    level: Literal["student", "graduate", "junior", "mid", "senior", "lead", "principal", "executive"]
    years: float
    primary: str = Field(description="The person's main specialization in a few words")
    summary: str = Field(description="3-4 honest sentences: where they stand and what they are ready for")
    strengths: list[Strength]
    transferable: list[str] = Field(description="Strengths that carry over to adjacent roles")
    growth_areas: list[str] = Field(description="Skills that would open the stretch tracks further")
    tracks: list[TrackPlan] = Field(description="3-6 tracks, core first")
    title_exclude: list[str] = Field(description="Title words that signal the wrong level for this person")


def plan_career(profile: Profile, country_name: str, llm: LLM | None = None) -> CareerPlan:
    from jobright import engine, heuristics
    if not engine.uses_ai():
        return heuristics.plan_career(profile, country_name)
    llm = llm or LLM()
    try:
        return _structured_plan(llm, profile, country_name)
    except Exception:
        if engine.current() != "local":
            raise
        return heuristics.plan_career(profile, country_name)


def _structured_plan(llm: LLM, profile: Profile, country_name: str) -> CareerPlan:
    return llm.structured(
        purpose="onboard_career",
        model=config.settings()["models"]["tailor"],
        instructions=config.prompt("onboard_career"),
        context="<master_profile>\n" + profile.for_model() + "</master_profile>",
        prompt=f"The person is job hunting in: {country_name}. Propose their career plan and role tracks.",
        schema=CareerPlan,
        effort="high",
        max_tokens=24000,
    )


def save_career(plan: CareerPlan) -> Path:
    config.CAREER_PATH.write_text(yaml.safe_dump(plan.model_dump(), sort_keys=False, allow_unicode=True, width=120))
    return config.CAREER_PATH


def load_career() -> CareerPlan | None:
    if not config.CAREER_PATH.exists():
        return None
    with config.CAREER_PATH.open() as f:
        return CareerPlan.model_validate(yaml.safe_load(f))


def searches_for(plan: CareerPlan, country: dict, cities: list[str], sites: list[str]) -> dict:
    """Site searches from the plan's keywords, for the chosen cities."""
    geo = country.get("linkedin_geo") or {}
    out: dict[str, list[dict]] = {}
    keywords = [k for t in plan.tracks for k in t.search_keywords[:3]]
    core = plan.tracks[0].search_keywords[:2] if plan.tracks else []
    if "linkedin" in sites:
        # Every keyword country-wide; the core track's also in each chosen city with a known geo id.
        country_geo = geo.get("country", 92000000)
        out["linkedin"] = [{"keywords": kw, "geo_id": country_geo, "label": country.get("name", "")} for kw in keywords]
        out["linkedin"] += [{"keywords": kw, "geo_id": geo[c], "label": c} for c in cities if c in geo for kw in core]
    city = (cities or [country.get("name", "")])[0]
    if "naukri" in sites:
        out["naukri"] = [{"keywords": kw.lower(), "location": (cities[i % len(cities)] if cities else "india").lower()}
                         for i, kw in enumerate(keywords)]
    if "reed" in sites:
        out["reed"] = [{"keywords": kw, "location": cities[i % len(cities)] if cities else ""} for i, kw in enumerate(keywords)]
    if "indeed" in sites:
        out["indeed"] = [{"keywords": kw, "location": cities[i % len(cities)] if cities else city} for i, kw in enumerate(keywords)]
    return out


def apply_plan(plan: CareerPlan, country_code: str, cities: list[str], sites: list[str],
               exclude_companies: list[str], remote_ok: bool = True) -> None:
    country = {"code": country_code, **config.countries().get(country_code, {})}
    aliases = [a for c in cities for a in (country.get("aliases") or {}).get(c, [])]
    locations = list(dict.fromkeys(cities + aliases + [country.get("name", "")] + (["Remote"] if remote_ok else [])))
    defaults = config._load(config.DEFAULTS_DIR / "settings.yaml")
    config.save_settings({
        "profile": {"country": country_code, "cities": cities, "sites": sites},
        "targets": {"title_exclude": list(dict.fromkeys(defaults["targets"]["title_exclude"] + plan.title_exclude)),
                    "locations": [l for l in locations if l]},
        "tracks": {t.name: {"label": t.label, "weight": t.weight, "min_fit": t.min_fit, "stretch": t.stretch,
                            "titles": t.titles} for t in plan.tracks},
        "searches": searches_for(plan, country, cities, sites),
        "ats_boards": country.get("ats_boards") or {},
        "exclude_companies": exclude_companies,
    })
